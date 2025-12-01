from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import io
import re

class DocGenService:
    def create_word_doc(self, content: str, title: str) -> io.BytesIO:
        """
        Converts Markdown-like content to a Word document.
        Returns a BytesIO object containing the .docx file.
        """
        doc = Document()
        
        # Title
        heading = doc.add_heading(title, 0)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Simple Markdown parsing
        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif re.match(r'^\d+\. ', line):
                # Ordered list
                text = re.sub(r'^\d+\. ', '', line)
                doc.add_paragraph(text, style='List Number')
            else:
                doc.add_paragraph(line)

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
